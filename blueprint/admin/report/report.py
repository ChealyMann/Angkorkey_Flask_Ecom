from datetime import datetime, time
from io import BytesIO
from zoneinfo import ZoneInfo

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, desc

from extensions import db
from models import Order, OrderItem, Product, ProductVariant

report_bp = Blueprint('report', __name__)


def get_status_label(status):
    status_map = {
        1: 'Pending',
        2: 'Confirmed',
        3: 'Shipped',
        4: 'Delivered',
        5: 'Canceled',
        6: 'Returned',
    }
    return status_map.get(status, 'Unknown')


def parse_single_date(date_str, default_date):
    try:
        return datetime.strptime(date_str, '%Y-%m-%d').date() if date_str else default_date
    except ValueError:
        return default_date


def parse_date_range(from_str, to_str, default_date):
    date_from = parse_single_date(from_str, default_date)
    date_to = parse_single_date(to_str, default_date)

    if date_from > date_to:
        date_from, date_to = date_to, date_from

    return date_from, date_to


def get_range_bounds(date_from, date_to):
    cambodia_tz = ZoneInfo("Asia/Phnom_Penh")
    start_dt = datetime.combine(date_from, time.min).replace(tzinfo=cambodia_tz)
    end_dt = datetime.combine(date_to, time.max).replace(tzinfo=cambodia_tz)
    return start_dt, end_dt


def apply_sheet_style_and_width(wb, border):
    for sheet in wb.worksheets:
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                cell.border = border

        for col_idx in range(1, sheet.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_length = 0

            for row_idx in range(1, sheet.max_row + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                value = '' if cell.value is None else str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)

            sheet.column_dimensions[col_letter].width = max_length + 3


def build_sales_report_data(filter_date):
    cambodia_tz = ZoneInfo("Asia/Phnom_Penh")

    start_day = datetime.combine(filter_date, time.min).replace(tzinfo=cambodia_tz)
    end_day = datetime.combine(filter_date, time.max).replace(tzinfo=cambodia_tz)
    month_start = datetime(filter_date.year, filter_date.month, 1, tzinfo=cambodia_tz)

    sales_statuses = [2, 3, 4]  # confirmed, shipped, delivered

    today_orders_query = Order.query.filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day
    )

    today_orders = today_orders_query.count()

    today_sales = db.session.query(
        func.coalesce(func.sum(Order.grand_total), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    month_sales = db.session.query(
        func.coalesce(func.sum(Order.grand_total), 0)
    ).filter(
        Order.created_at >= month_start,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    avg_order_value = db.session.query(
        func.coalesce(func.avg(Order.grand_total), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    pending_orders = today_orders_query.filter(Order.status == 1).count()
    confirmed_orders = today_orders_query.filter(Order.status == 2).count()
    shipped_orders = today_orders_query.filter(Order.status == 3).count()
    delivered_orders = today_orders_query.filter(Order.status == 4).count()
    canceled_orders = today_orders_query.filter(Order.status == 5).count()
    returned_orders = today_orders_query.filter(Order.status == 6).count()

    gross_sales = db.session.query(
        func.coalesce(func.sum(Order.sub_total), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    total_discount = db.session.query(
        func.coalesce(func.sum(Order.discount), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    total_shipping = db.session.query(
        func.coalesce(func.sum(Order.shipping_fee), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    net_sales = db.session.query(
        func.coalesce(func.sum(Order.grand_total), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    cod_sales = db.session.query(
        func.coalesce(func.sum(Order.grand_total), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses),
        Order.payment_method == 'cash_on_delivery'
    ).scalar() or 0

    online_sales = db.session.query(
        func.coalesce(func.sum(Order.grand_total), 0)
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses),
        Order.payment_method != 'cash_on_delivery'
    ).scalar() or 0

    paid_orders = today_orders_query.filter(Order.payment_status == 'paid').count()
    pending_payment_orders = today_orders_query.filter(Order.payment_status == 'pending').count()

    sold_units = db.session.query(
        func.coalesce(func.sum(OrderItem.qty), 0)
    ).join(Order, Order.id == OrderItem.order_id).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    profit_rows = db.session.query(
        OrderItem.qty,
        OrderItem.unit_price,
        ProductVariant.purchase_cost
    ).join(
        Order, Order.id == OrderItem.order_id
    ).outerjoin(
        ProductVariant, ProductVariant.id == OrderItem.variant_id
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).all()

    estimated_profit = 0
    for row in profit_rows:
        if row.purchase_cost is not None:
            estimated_profit += (float(row.unit_price) - float(row.purchase_cost)) * row.qty

    top_products_raw = db.session.query(
        Product.id.label('product_id'),
        Product.name.label('name'),
        func.coalesce(func.sum(OrderItem.qty), 0).label('qty_sold'),
        func.coalesce(func.sum(OrderItem.sub_total), 0).label('sales')
    ).join(
        OrderItem, Product.id == OrderItem.product_id
    ).join(
        Order, Order.id == OrderItem.order_id
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).group_by(
        Product.id, Product.name
    ).order_by(
        desc('qty_sold')
    ).limit(5).all()

    top_products = [
        {
            'product_id': row.product_id,
            'name': row.name,
            'qty_sold': int(row.qty_sold or 0),
            'sales': float(row.sales or 0)
        }
        for row in top_products_raw
    ]

    top_variants_raw = db.session.query(
        Product.name.label('product_name'),
        ProductVariant.id.label('variant_id'),
        ProductVariant.product_id.label('product_id'),
        ProductVariant.color,
        ProductVariant.type,
        func.coalesce(func.sum(OrderItem.qty), 0).label('qty_sold'),
        func.coalesce(func.sum(OrderItem.sub_total), 0).label('sales')
    ).join(
        OrderItem, ProductVariant.id == OrderItem.variant_id
    ).join(
        Product, Product.id == OrderItem.product_id
    ).join(
        Order, Order.id == OrderItem.order_id
    ).filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day,
        Order.status.in_(sales_statuses)
    ).group_by(
        Product.name, ProductVariant.id, ProductVariant.product_id, ProductVariant.color, ProductVariant.type
    ).order_by(
        desc('qty_sold')
    ).limit(5).all()

    top_variants = [
        {
            'variant_id': row.variant_id,
            'product_id': row.product_id,
            'product_name': row.product_name,
            'color': row.color or '-',
            'type': row.type or '-',
            'qty_sold': int(row.qty_sold or 0),
            'sales': float(row.sales or 0)
        }
        for row in top_variants_raw
    ]

    recent_orders_raw = Order.query.filter(
        Order.created_at >= start_day,
        Order.created_at <= end_day
    ).order_by(
        Order.created_at.desc()
    ).limit(10).all()

    recent_orders = []
    for order in recent_orders_raw:
        recent_orders.append({
            'id': order.id,
            'invoice_no': order.invoice_no,
            'created_at': order.created_at,
            'payment_method': order.payment_method,
            'payment_status': order.payment_status,
            'status': order.status,
            'status_label': get_status_label(order.status),
            'grand_total': float(order.grand_total or 0)
        })

    low_stock_raw = db.session.query(
        Product.id.label('product_id'),
        Product.name.label('product_name'),
        ProductVariant.id.label('variant_id'),
        ProductVariant.sku,
        ProductVariant.color,
        ProductVariant.type,
        ProductVariant.price,
        ProductVariant.discount_price,
        ProductVariant.purchase_cost,
        ProductVariant.physical_stock,
        ProductVariant.reserved_stock
    ).join(
        Product, Product.id == ProductVariant.product_id
    ).all()

    low_stock = []
    for row in low_stock_raw:
        physical_stock = row.physical_stock or 0
        reserved_stock = row.reserved_stock or 0
        available_stock = physical_stock - reserved_stock

        if available_stock <= 5:
            low_stock.append({
                'product_id': row.product_id,
                'variant_id': row.variant_id,
                'sku': row.sku or '-',
                'product_name': row.product_name,
                'color': row.color or '-',
                'type': row.type or '-',
                'price': float(row.discount_price if row.discount_price is not None else (row.price or 0)),
                'purchase_cost': float(row.purchase_cost or 0),
                'physical_stock': physical_stock,
                'reserved_stock': reserved_stock,
                'available_stock': max(available_stock, 0)
            })

    low_stock = sorted(low_stock, key=lambda x: x['available_stock'])[:10]

    report = {
        'today_sales': float(today_sales),
        'month_sales': float(month_sales),
        'today_orders': int(today_orders),
        'avg_order_value': float(avg_order_value),
        'pending_orders': int(pending_orders),
        'confirmed_orders': int(confirmed_orders),
        'shipped_orders': int(shipped_orders),
        'delivered_orders': int(delivered_orders),
        'canceled_orders': int(canceled_orders),
        'returned_orders': int(returned_orders),
        'gross_sales': float(gross_sales),
        'total_discount': float(total_discount),
        'total_shipping': float(total_shipping),
        'net_sales': float(net_sales),
        'cod_sales': float(cod_sales),
        'online_sales': float(online_sales),
        'paid_orders': int(paid_orders),
        'pending_payment_orders': int(pending_payment_orders),
        'estimated_profit': float(estimated_profit),
        'sold_units': int(sold_units),
        'top_products': top_products,
        'top_variants': top_variants,
        'recent_orders': recent_orders,
        'low_stock': low_stock,
    }

    return report


def build_low_stock_report_data():
    rows = db.session.query(
        Product.id.label('product_id'),
        Product.name.label('product_name'),
        ProductVariant.id.label('variant_id'),
        ProductVariant.sku,
        ProductVariant.color,
        ProductVariant.type,
        ProductVariant.price,
        ProductVariant.discount_price,
        ProductVariant.purchase_cost,
        ProductVariant.physical_stock,
        ProductVariant.reserved_stock
    ).join(
        Product, Product.id == ProductVariant.product_id
    ).order_by(
        Product.name.asc(),
        ProductVariant.color.asc(),
        ProductVariant.type.asc()
    ).all()

    items = []
    total_items = 0
    total_available_units = 0

    for row in rows:
        physical_stock = row.physical_stock or 0
        reserved_stock = row.reserved_stock or 0
        available_stock = physical_stock - reserved_stock

        if available_stock <= 5:
            total_items += 1
            total_available_units += max(available_stock, 0)

            items.append({
                'product_id': row.product_id,
                'variant_id': row.variant_id,
                'product_name': row.product_name,
                'sku': row.sku or '-',
                'color': row.color or '-',
                'type': row.type or '-',
                'price': row.price or 0,
                'purchase_cost': float(row.purchase_cost or 0),
                'physical_stock': physical_stock,
                'reserved_stock': reserved_stock,
                'available_stock': max(available_stock, 0),
                'stock_status': 'Out of Stock' if available_stock <= 0 else 'Low Stock'
            })

    items = sorted(items, key=lambda x: (x['available_stock'], x['product_name']))

    return {
        'total_items': total_items,
        'total_available_units': total_available_units,
        'items': items
    }


def build_profit_loss_report_data(date_from, date_to):
    start_dt, end_dt = get_range_bounds(date_from, date_to)
    sales_statuses = [2, 3, 4]

    orders_count = Order.query.filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).count()

    gross_sales = db.session.query(
        func.coalesce(func.sum(Order.sub_total), 0)
    ).filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    total_discount = db.session.query(
        func.coalesce(func.sum(Order.discount), 0)
    ).filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    total_shipping = db.session.query(
        func.coalesce(func.sum(Order.shipping_fee), 0)
    ).filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    net_sales = db.session.query(
        func.coalesce(func.sum(Order.grand_total), 0)
    ).filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    sold_units = db.session.query(
        func.coalesce(func.sum(OrderItem.qty), 0)
    ).join(
        Order, Order.id == OrderItem.order_id
    ).filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).scalar() or 0

    item_rows = db.session.query(
        Order.invoice_no,
        Order.created_at,
        Product.name.label('product_name'),
        ProductVariant.sku,
        ProductVariant.color,
        ProductVariant.type,
        OrderItem.qty,
        OrderItem.unit_price,
        OrderItem.sub_total,
        Product.cost.label('product_cost'),
        ProductVariant.purchase_cost
    ).join(
        Order, Order.id == OrderItem.order_id
    ).join(
        Product, Product.id == OrderItem.product_id
    ).outerjoin(
        ProductVariant, ProductVariant.id == OrderItem.variant_id
    ).filter(
        Order.created_at >= start_dt,
        Order.created_at <= end_dt,
        Order.status.in_(sales_statuses)
    ).order_by(
        Order.created_at.desc()
    ).all()

    detail_rows = []
    total_cogs = 0.0
    total_profit = 0.0

    for row in item_rows:
        unit_cost = row.purchase_cost if row.purchase_cost is not None else row.product_cost
        unit_cost = float(unit_cost or 0)
        qty = int(row.qty or 0)
        unit_price = float(row.unit_price or 0)
        sales_amount = float(row.sub_total or 0)
        cogs = unit_cost * qty
        profit = sales_amount - cogs

        total_cogs += cogs
        total_profit += profit

        detail_rows.append({
            'invoice_no': row.invoice_no,
            'created_at': row.created_at,
            'product_name': row.product_name,
            'sku': row.sku or '-',
            'color': row.color or '-',
            'type': row.type or '-',
            'qty': qty,
            'unit_price': unit_price,
            'unit_cost': unit_cost,
            'sales_amount': sales_amount,
            'cogs': cogs,
            'profit': profit
        })

    return {
        'date_from': date_from.strftime('%Y-%m-%d'),
        'date_to': date_to.strftime('%Y-%m-%d'),
        'orders_count': orders_count,
        'sold_units': int(sold_units or 0),
        'gross_sales': float(gross_sales or 0),
        'total_discount': float(total_discount or 0),
        'total_shipping': float(total_shipping or 0),
        'net_sales': float(net_sales or 0),
        'total_cogs': float(total_cogs),
        'total_profit': float(total_profit),
        'net_result': float(total_profit),
        'detail_rows': detail_rows
    }


@report_bp.route('/admin/report/sales')
def sales_report():
    cambodia_tz = ZoneInfo("Asia/Phnom_Penh")
    today_kh = datetime.now(cambodia_tz).date()

    selected_date = request.args.get('report_date', '').strip()
    filter_date = parse_single_date(selected_date, today_kh)

    report = build_sales_report_data(filter_date)

    return render_template(
        'backend/admin/pages/report/sales_report.html',
        report=report,
        selected_date=filter_date.strftime('%Y-%m-%d')
    )


@report_bp.route('/admin/report/sales/export')
def export_sales_report():
    cambodia_tz = ZoneInfo("Asia/Phnom_Penh")
    today_kh = datetime.now(cambodia_tz).date()

    selected_date = request.args.get('report_date', '').strip()
    filter_date = parse_single_date(selected_date, today_kh)

    report = build_sales_report_data(filter_date)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    title_fill = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
    header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    section_fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells('A1:B1')
    ws['A1'] = 'Sales Report'
    ws['A1'].fill = title_fill
    ws['A1'].font = Font(color='FFFFFF', bold=True, size=16)
    ws['A1'].alignment = center

    ws['A2'] = 'Report Date'
    ws['B2'] = filter_date.strftime('%Y-%m-%d')
    ws['A2'].font = bold_font
    ws['B2'].font = bold_font

    row = 4
    ws[f'A{row}'] = 'Summary'
    ws[f'A{row}'].fill = section_fill
    ws[f'A{row}'].font = bold_font

    row += 1
    ws[f'A{row}'] = 'Metric'
    ws[f'B{row}'] = 'Value'
    for cell in ws[row]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    summary_rows = [
        ('Today Sales', report['today_sales']),
        ('This Month Sales', report['month_sales']),
        ('Today Orders', report['today_orders']),
        ('Average Order Value', report['avg_order_value']),
        ('Gross Sales', report['gross_sales']),
        ('Discount', report['total_discount']),
        ('Shipping Fee', report['total_shipping']),
        ('Net Sales', report['net_sales']),
        ('Cash on Delivery Sales', report['cod_sales']),
        ('Online Payment Sales', report['online_sales']),
        ('Paid Orders', report['paid_orders']),
        ('Pending Payment Orders', report['pending_payment_orders']),
        ('Estimated Profit', report['estimated_profit']),
        ('Sold Units', report['sold_units']),
    ]

    for metric, value in summary_rows:
        row += 1
        ws[f'A{row}'] = metric
        ws[f'B{row}'] = value
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border

    row += 2
    ws[f'A{row}'] = 'Order Status Summary'
    ws[f'A{row}'].fill = section_fill
    ws[f'A{row}'].font = bold_font

    row += 1
    ws[f'A{row}'] = 'Status'
    ws[f'B{row}'] = 'Count'
    for cell in ws[row]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    status_rows = [
        ('Pending', report['pending_orders']),
        ('Confirmed', report['confirmed_orders']),
        ('Shipped', report['shipped_orders']),
        ('Delivered', report['delivered_orders']),
        ('Canceled', report['canceled_orders']),
        ('Returned', report['returned_orders']),
    ]

    for label, count in status_rows:
        row += 1
        ws[f'A{row}'] = label
        ws[f'B{row}'] = count
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border

    top_products_ws = wb.create_sheet('Top Products')
    top_products_ws.append(['Product', 'Qty Sold', 'Sales'])
    for cell in top_products_ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    for item in report['top_products']:
        top_products_ws.append([
            item['name'],
            item['qty_sold'],
            item['sales']
        ])

    top_variants_ws = wb.create_sheet('Top Variants')
    top_variants_ws.append(['Product', 'Color', 'Type', 'Qty Sold', 'Sales'])
    for cell in top_variants_ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    for item in report['top_variants']:
        top_variants_ws.append([
            item['product_name'],
            item['color'],
            item['type'],
            item['qty_sold'],
            item['sales']
        ])

    recent_orders_ws = wb.create_sheet('Recent Orders')
    recent_orders_ws.append(['Invoice', 'Date', 'Payment Method', 'Payment Status', 'Status', 'Total'])
    for cell in recent_orders_ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    for item in report['recent_orders']:
        recent_orders_ws.append([
            item['invoice_no'],
            item['created_at'].strftime('%Y-%m-%d %H:%M') if item['created_at'] else '',
            item['payment_method'],
            item['payment_status'],
            item['status_label'],
            item['grand_total']
        ])

    low_stock_ws = wb.create_sheet('Low Stock')
    low_stock_ws.append(['Product', 'SKU', 'Color', 'Type', 'Physical Stock', 'Reserved Stock', 'Available Stock'])
    for cell in low_stock_ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    for item in report['low_stock']:
        low_stock_ws.append([
            item['product_name'],
            item['sku'],
            item['color'],
            item['type'],
            item['physical_stock'],
            item['reserved_stock'],
            item['available_stock']
        ])

    apply_sheet_style_and_width(wb, border)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"sales_report_{filter_date.strftime('%Y-%m-%d')}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@report_bp.route('/admin/report/low-stock')
def low_stock_report():
    report = build_low_stock_report_data()

    return render_template(
        'backend/admin/pages/report/low_stock_report.html',
        report=report
    )


@report_bp.route('/admin/report/low-stock/export')
def export_low_stock_report():
    report = build_low_stock_report_data()

    wb = Workbook()
    ws = wb.active
    ws.title = "Low Stock Report"

    title_fill = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
    header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 8 columns only now: A to H
    ws.merge_cells('A1:H1')
    ws['A1'] = 'Low Stock Report'
    ws['A1'].fill = title_fill
    ws['A1'].font = Font(color='FFFFFF', bold=True, size=16)
    ws['A1'].alignment = center

    ws['A2'] = 'Total Low Stock Items'
    ws['B2'] = report['total_items']
    ws['A3'] = 'Total Available Units'
    ws['B3'] = report['total_available_units']

    ws['A2'].font = bold_font
    ws['A3'].font = bold_font

    row = 5

    headers = [
        'Product',
        'SKU',
        'Color',
        'Type',
        'Physical Stock',
        'Reserved Stock',
        'Available Stock',
        'Stock Status'
    ]

    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=header)
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    for item in report['items']:
        row += 1

        ws.cell(row=row, column=1, value=item['product_name'])
        ws.cell(row=row, column=2, value=item['sku'])
        ws.cell(row=row, column=3, value=item['color'])
        ws.cell(row=row, column=4, value=item['type'])
        ws.cell(row=row, column=5, value=item['physical_stock'])
        ws.cell(row=row, column=6, value=item['reserved_stock'])
        ws.cell(row=row, column=7, value=item['available_stock'])
        ws.cell(row=row, column=8, value=item['stock_status'])

        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = center

    apply_sheet_style_and_width(wb, border)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "low_stock_report_.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@report_bp.route('/admin/report/profit-loss')
def profit_loss_report():
    cambodia_tz = ZoneInfo("Asia/Phnom_Penh")
    today_kh = datetime.now(cambodia_tz).date()

    from_date_str = request.args.get('from_date', '').strip()
    to_date_str = request.args.get('to_date', '').strip()

    date_from, date_to = parse_date_range(from_date_str, to_date_str, today_kh)
    report = build_profit_loss_report_data(date_from, date_to)

    return render_template(
        'backend/admin/pages/report/profit_loss_report.html',
        report=report,
        from_date=report['date_from'],
        to_date=report['date_to']
    )


@report_bp.route('/admin/report/profit-loss/export')
def export_profit_loss_report():
    cambodia_tz = ZoneInfo("Asia/Phnom_Penh")
    today_kh = datetime.now(cambodia_tz).date()

    from_date_str = request.args.get('from_date', '').strip()
    to_date_str = request.args.get('to_date', '').strip()

    date_from, date_to = parse_date_range(from_date_str, to_date_str, today_kh)
    report = build_profit_loss_report_data(date_from, date_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "Profit Loss Summary"

    title_fill = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
    header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    section_fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells('A1:B1')
    ws['A1'] = 'Profit & Loss Report'
    ws['A1'].fill = title_fill
    ws['A1'].font = Font(color='FFFFFF', bold=True, size=16)
    ws['A1'].alignment = center

    ws['A2'] = 'From Date'
    ws['B2'] = report['date_from']
    ws['A3'] = 'To Date'
    ws['B3'] = report['date_to']
    ws['A2'].font = bold_font
    ws['A3'].font = bold_font

    row = 5
    ws[f'A{row}'] = 'Summary'
    ws[f'A{row}'].fill = section_fill
    ws[f'A{row}'].font = bold_font

    row += 1
    ws[f'A{row}'] = 'Metric'
    ws[f'B{row}'] = 'Value'
    for cell in ws[row]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    summary_rows = [
        ('Orders Count', report['orders_count']),
        ('Sold Units', report['sold_units']),
        ('Gross Sales', report['gross_sales']),
        ('Discount', report['total_discount']),
        ('Shipping Fee', report['total_shipping']),
        ('Net Sales', report['net_sales']),
        ('Cost of Goods Sold', report['total_cogs']),
        ('Total Profit', report['total_profit']),
        ('Net Result', report['net_result']),
    ]

    for metric, value in summary_rows:
        row += 1
        ws[f'A{row}'] = metric
        ws[f'B{row}'] = value
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border

    details_ws = wb.create_sheet('Profit Loss Details')
    detail_headers = [
        'Invoice', 'Date', 'Product', 'SKU', 'Color', 'Type',
        'Qty', 'Unit Price', 'Unit Cost', 'Sales Amount', 'COGS', 'Profit'
    ]
    details_ws.append(detail_headers)

    for cell in details_ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    for item in report['detail_rows']:
        details_ws.append([
            item['invoice_no'],
            item['created_at'].strftime('%Y-%m-%d %H:%M') if item['created_at'] else '',
            item['product_name'],
            item['sku'],
            item['color'],
            item['type'],
            item['qty'],
            item['unit_price'],
            item['unit_cost'],
            item['sales_amount'],
            item['cogs'],
            item['profit'],
        ])

    apply_sheet_style_and_width(wb, border)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"profit_loss_{report['date_from']}_to_{report['date_to']}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )